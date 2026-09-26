"""Inspect workbook formulas and error cells; this script does not calculate."""

import argparse
import json
from pathlib import Path
import sys
import zipfile


def inspect(path: Path) -> dict:
    from openpyxl import load_workbook

    formulas = load_workbook(path, read_only=True, data_only=False)
    try:
        values = load_workbook(path, read_only=True, data_only=True)
        try:
            sheets, errors, missing = [], [], []
            formula_count = 0
            for sheet in formulas:
                filled = 0
                count = 0
                for row, cached_row in zip(sheet.iter_rows(), values[sheet.title].iter_rows()):
                    for cell, cached in zip(row, cached_row):
                        if cell.value is not None:
                            filled += 1
                        if cell.data_type == "f":
                            count += 1
                            if cached.value is None:
                                missing.append({"sheet": sheet.title, "cell": cell.coordinate})
                        error = cell if cell.data_type == "e" else cached if cached.data_type == "e" else None
                        if error is not None:
                            errors.append({"sheet": sheet.title, "cell": cell.coordinate, "error": error.value})
                formula_count += count
                sheets.append({"name": sheet.title, "filled_cells": filled, "formula_count": count})
            return {
                "sheets": sheets, "formula_count": formula_count, "errors": errors,
                "missing_cached_results": missing, "calculated_by_this_script": False,
                "not_checked": ["cache_freshness", "formula_semantics", "visual_layout"],
            }
        finally:
            values.close()
    finally:
        formulas.close()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("file", type=Path)
    parser.add_argument("--require-cache", action="store_true")
    args = parser.parse_args()
    try:
        result = inspect(args.file)
    except ImportError:
        print("Missing dependency: install openpyxl in the task Python environment.", file=sys.stderr)
        return 2
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
        print(f"Cannot inspect workbook: {exc}", file=sys.stderr)
        return 2
    print(json.dumps(result, ensure_ascii=True, indent=2))
    return int(bool(result["errors"] or (args.require_cache and result["missing_cached_results"])))


if __name__ == "__main__":
    raise SystemExit(main())

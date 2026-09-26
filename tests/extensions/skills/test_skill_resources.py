import json
import re
import runpy
import shutil
import subprocess
import sys

import pytest

from reuleauxcoder.extensions.skills.discovery import BUILTIN_SKILLS_DIR
from reuleauxcoder.extensions.skills.service import SkillsService


def helper(skill, script):
    return BUILTIN_SKILLS_DIR / skill / "scripts" / script


def command(script, *args, cwd, code=0):
    result = subprocess.run(
        [sys.executable, str(script), *map(str, args)], cwd=cwd,
        capture_output=True, text=True, encoding="utf-8", timeout=30,
    )
    assert result.returncode == code, result.stdout + result.stderr
    return result


def test_every_builtin_is_self_contained_after_copy(tmp_path):
    validate = runpy.run_path(str(helper("skill-creator", "validate_skill.py")))["validate"]
    for source in BUILTIN_SKILLS_DIR.iterdir():
        if not (source / "SKILL.md").is_file():
            continue
        target = tmp_path / source.name
        shutil.copytree(source, target, ignore=shutil.ignore_patterns("__pycache__"))
        assert validate(target) == [], source.name
        for script in (target / "scripts").glob("*.py"):
            command(script, "--help", cwd=tmp_path)


def test_validator_rejects_missing_and_external_resources(tmp_path):
    root = tmp_path / "example"
    root.mkdir()
    entry = root / "SKILL.md"
    entry.write_text(
        "---\nname: example\ndescription: Example\n---\n"
        "[Missing](references/missing.md)\n[Outside](../outside.md)\n", encoding="utf-8",
    )
    result = command(helper("skill-creator", "validate_skill.py"), root, cwd=tmp_path, code=1)
    assert "Missing resource" in result.stderr
    assert "escapes skill" in result.stderr


def test_install_preserves_resources_and_never_overwrites(tmp_path):
    source = tmp_path / "source" / "example"
    (source / "references").mkdir(parents=True)
    (source / "SKILL.md").write_text(
        "---\nname: example\ndescription: Example\n---\n[Detail](references/detail.md)\n", encoding="utf-8",
    )
    (source / "references/detail.md").write_text("Useful detail", encoding="utf-8")
    workspace = tmp_path / "workspace"
    root = workspace / ".rcoder/skills"
    script = helper("skill-installer", "install_local.py")
    command(script, source, root, cwd=tmp_path)
    installed = root / "example"
    assert (installed / "references/detail.md").read_text() == "Useful detail"
    loaded = SkillsService(workspace_dir=workspace, home_dir=tmp_path / "home").reload()
    assert next(skill for skill in loaded.active_skills if skill.name == "example").scope == "project"
    (installed / "references/detail.md").write_text("User edit", encoding="utf-8")
    command(script, source, root, cwd=tmp_path, code=1)
    assert (installed / "references/detail.md").read_text() == "User edit"
    assert not list(root.glob(".*"))


def test_install_rejects_path_names_and_symlinks(tmp_path):
    install = runpy.run_path(str(helper("skill-installer", "install_local.py")))["install"]
    source = tmp_path / "example"
    source.mkdir()
    entry = source / "SKILL.md"
    entry.write_text("---\nname: ../escape\ndescription: Example\n---\nBody", encoding="utf-8")
    with pytest.raises(ValueError, match="name"):
        install(source, tmp_path / "target")
    assert not (tmp_path / "target").exists()
    entry.write_text("---\nname: example\ndescription: Example\n---\nBody", encoding="utf-8")
    try:
        (source / "outside").symlink_to(entry)
    except OSError:
        pytest.skip("Symlinks unavailable on this host")
    with pytest.raises(ValueError, match="Symlinks"):
        install(source, tmp_path / "target")
    assert not (tmp_path / "target").exists()


def test_pptx_inspector_reads_real_slides_and_detects_geometry(tmp_path):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    deck = pptx.Presentation()
    for title in ("First", "Second"):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = title
    file = tmp_path / "演示 with spaces.pptx"
    deck.save(file)
    original = file.read_bytes()
    script = helper("pptx", "inspect_deck.py")
    result = json.loads(command(script, file, cwd=tmp_path).stdout)
    assert result["slide_count"] == 2
    assert [slide["shapes"][0]["text"] for slide in result["slides"]] == ["First", "Second"]
    assert file.read_bytes() == original
    deck.slides[0].shapes[0].left = -Inches(1)
    deck.save(file)
    result = json.loads(command(script, file, cwd=tmp_path, code=1).stdout)
    assert result["warnings"][0]["issue"] == "outside_canvas"


def test_docx_inspector_preserves_comments_and_revisions(tmp_path):
    docx = pytest.importorskip("docx")
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = docx.Document()
    paragraph = doc.add_paragraph("Original text")
    doc.add_comment(paragraph.runs, text="Review this", author="Reviewer", initials="R")
    inserted = OxmlElement("w:ins")
    inserted.set(qn("w:id"), "1")
    inserted.set(qn("w:author"), "Reviewer")
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "Inserted text"
    run.append(text)
    inserted.append(run)
    paragraph._p.append(inserted)
    doc.add_table(rows=1, cols=2)
    file = tmp_path / "报告 with spaces.docx"
    doc.save(file)
    original = file.read_bytes()
    result = json.loads(command(helper("docx", "inspect_doc.py"), file, cwd=tmp_path).stdout)
    assert result["table_count"] == 1
    assert result["revisions"] == {"insertions": 1, "deletions": 0, "comment_anchors": 1}
    assert result["comments"][0]["text"] == "Review this"
    assert "Inserted text" in result["paragraphs"][0]["text"]
    assert file.read_bytes() == original


def test_workbook_check_distinguishes_uncalculated_formulas_and_errors(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append([2, 3, "=A1+B1", "00123"])
    file = tmp_path / "数据 with spaces.xlsx"
    book.save(file)
    script = helper("xlsx", "check_workbook.py")
    result = json.loads(command(script, file, cwd=tmp_path).stdout)
    assert result["formula_count"] == 1
    assert result["missing_cached_results"] == [{"sheet": "Sheet", "cell": "C1"}]
    assert not result["calculated_by_this_script"]
    command(script, file, "--require-cache", cwd=tmp_path, code=1)
    sheet["E1"] = "#DIV/0!"
    book.save(file)
    result = json.loads(command(script, file, cwd=tmp_path, code=1).stdout)
    assert result["errors"] == [{"sheet": "Sheet", "cell": "E1", "error": "#DIV/0!"}]
    assert openpyxl.load_workbook(file).active["D1"].value == "00123"


def test_pdf_processing_and_rendering_real_content(tmp_path):
    pytest.importorskip("reportlab")
    pytest.importorskip("pypdf")
    pytest.importorskip("pymupdf")
    from reportlab.pdfgen import canvas
    from pypdf import PdfReader
    from PIL import Image

    source = tmp_path / "文档 with spaces.pdf"
    pdf = canvas.Canvas(str(source))
    for word in ("First page", "Second page"):
        pdf.drawString(72, 700, word)
        pdf.showPage()
    pdf.save()
    original = source.read_bytes()
    script = helper("pdf", "pdf_tools.py")
    info = json.loads(command(script, "info", source, cwd=tmp_path).stdout)
    assert info["pages"] == 2
    assert all(info["text_characters"])
    extracted = tmp_path / "extracted.txt"
    command(script, "extract", source, extracted, cwd=tmp_path)
    assert "Second page" in extracted.read_text(encoding="utf-8")
    command(script, "extract", source, extracted, cwd=tmp_path, code=2)
    selected = tmp_path / "selected.pdf"
    command(script, "select", source, selected, "--pages", "2,1", cwd=tmp_path)
    assert "Second page" in PdfReader(selected).pages[0].extract_text()
    merged = tmp_path / "merged.pdf"
    command(script, "merge", merged, source, selected, cwd=tmp_path)
    assert len(PdfReader(merged).pages) == 4
    preview = tmp_path / "preview"
    command(script, "render", source, preview, "--pages", "2", cwd=tmp_path)
    images = list(preview.glob("*.png"))
    assert len(images) == 1
    with Image.open(images[0]) as image:
        assert image.width > 900 and image.height > 1200
        assert image.convert("L").getextrema()[0] < 100  # Text was actually rendered.
    command(script, "select", source, tmp_path / "invalid.pdf", "--pages", "0", cwd=tmp_path, code=2)
    assert not (tmp_path / "invalid.pdf").exists()
    assert source.read_bytes() == original


@pytest.mark.parametrize("skill,reference,dependency,output", [
    ("pptx", "creation.md", "pptx", "presentation.pptx"),
    ("docx", "authoring.md", "docx", "report.docx"),
    ("xlsx", "workbooks.md", "openpyxl", "analysis.xlsx"),
    ("pdf", "creation.md", "reportlab", "report.pdf"),
])
def test_documented_creation_example_runs_outside_checkout(tmp_path, skill, reference, dependency, output):
    pytest.importorskip(dependency)
    text = (BUILTIN_SKILLS_DIR / skill / "references" / reference).read_text(encoding="utf-8")
    example = re.search(r"```python\n(.*?)```", text, re.DOTALL).group(1)
    script = tmp_path / "example.py"
    script.write_text(example, encoding="utf-8")
    command(script, cwd=tmp_path)
    assert (tmp_path / output).stat().st_size > 100
    result = subprocess.run([sys.executable, str(script)], cwd=tmp_path, capture_output=True, timeout=30)
    assert result.returncode != 0  # Examples also preserve existing inputs/outputs.
